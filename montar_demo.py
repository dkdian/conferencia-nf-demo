# -*- coding: utf-8 -*-
"""
Monta tudo o que a demonstração precisa e deixa no ponto de partida:

  1. "Planilha Demo.xlsm": a aba de cadastro (Lista) + a aba Conferência NF com os 3 botões;
  2. o "Drive" fictício (drive_ficticio/) com os dois relatórios de fluxo de caixa da semana;
  3. o ERP fictício com todos os lançamentos em aberto e o chat vazio;
  4. a rotina de terça, que baixa os relatórios, cria a aba da semana e avisa o gestor no chat.

    python montar_demo.py                 faz tudo (inclusive a rotina de terça)
    python montar_demo.py --sem-rotina    deixa a planilha SEM a aba da semana (rode rotina_terca.py depois)

Rode de novo sempre que quiser recomeçar do zero. Precisa do Excel instalado.
"""

from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl
import pythoncom
import win32com.client as win32
from openpyxl.styles import Alignment, Font, PatternFill

import config as cfg

sys.path.insert(0, str(cfg.BASE / "dados"))
import dados_demo  # noqa: E402
import instalar_botao  # noqa: E402
from erp_ficticio import garantir  # noqa: E402

XLSX_TEMP = cfg.BASE / "_planilha_temp.xlsx"
DRIVE = cfg.BASE / "drive_ficticio"


def _estilo_cabecalho(ws, altura=32):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = altura


def criar_xlsx():
    """A planilha começa só com o cadastro: as abas da semana são criadas pela rotina de terça."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lista"
    ws.append(["Conta bancária", "Forma de pagamento"])
    for i in range(max(len(dados_demo.CONTAS), len(dados_demo.FORMAS))):
        ws.append([dados_demo.CONTAS[i] if i < len(dados_demo.CONTAS) else None,
                   dados_demo.FORMAS[i] if i < len(dados_demo.FORMAS) else None])
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 22
    _estilo_cabecalho(ws)
    wb.save(XLSX_TEMP)


def criar_drive():
    """Os dois relatórios que o cliente deixa no Drive, com os nomes de coluna e os valores do jeito do cliente."""
    shutil.rmtree(DRIVE, ignore_errors=True)
    DRIVE.mkdir()
    for rel in dados_demo.RELATORIOS:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Movimentos"
        ws.append(dados_demo.COLUNAS_BRUTAS)
        for l in dados_demo.linhas_brutas(rel):
            ws.append([l[c] for c in dados_demo.COLUNAS_BRUTAS])
        for i, w in enumerate([12, 12, 38, 10, 12, 14, 16, 18, 12, 28, 16, 10], 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        for linha in ws.iter_rows(min_row=2):
            linha[0].number_format = linha[1].number_format = "dd/mm/yyyy"
            linha[4].number_format = "#,##0.00"
        _estilo_cabecalho(ws)
        wb.save(DRIVE / ("Fluxo de caixa - %s.xlsx" % rel))


def xlsx_para_xlsm():
    """O Excel faz a conversão: é o jeito de ter um .xlsm sem macros de terceiros."""
    pythoncom.CoInitialize()
    excel = win32.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        if cfg.PLANILHA.exists():
            cfg.PLANILHA.unlink()
        wb = excel.Workbooks.Open(str(XLSX_TEMP))
        wb.SaveAs(str(cfg.PLANILHA), 52)      # xlOpenXMLWorkbookMacroEnabled
        wb.Close(SaveChanges=False)
    finally:
        excel.Quit()
        pythoncom.CoUninitialize()
    XLSX_TEMP.unlink(missing_ok=True)


def limpar_metadados():
    """O Excel grava o nome da conta de quem salvou no arquivo; a demonstração é pública, então vira "demo"."""
    import re
    import zipfile
    origem = cfg.PLANILHA
    temp = origem.with_suffix(".tmp")
    with zipfile.ZipFile(origem) as entrada, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as saida:
        for item in entrada.infolist():
            dados = entrada.read(item.filename)
            if item.filename == "docProps/core.xml":
                texto = dados.decode("utf-8")
                texto = re.sub(r"<dc:creator>.*?</dc:creator>", "<dc:creator>demo</dc:creator>", texto)
                texto = re.sub(r"<cp:lastModifiedBy>.*?</cp:lastModifiedBy>", "<cp:lastModifiedBy>demo</cp:lastModifiedBy>", texto)
                dados = texto.encode("utf-8")
            saida.writestr(item, dados)
    temp.replace(origem)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--sem-rotina", action="store_true", help="não roda a rotina de terça no fim")
    args = ap.parse_args()

    print("1/4  Criando a planilha com os botões...")
    criar_xlsx()
    xlsx_para_xlsm()
    if instalar_botao.main([]) != 0:
        return 1
    print("2/4  Criando o Drive fictício com os relatórios da semana...")
    criar_drive()
    print("3/4  Reiniciando o ERP fictício e limpando execuções anteriores...")
    garantir.reiniciar()
    for pasta in (cfg.SAIDA, cfg.BACKUP, cfg.BASE / "recebidos"):
        shutil.rmtree(pasta, ignore_errors=True)
    cfg.ARQ_INSERCAO.unlink(missing_ok=True)

    if args.sem_rotina:
        limpar_metadados()
        print("\nPronto: %s (sem a aba da semana). Rode: python rotina_terca.py" % cfg.PLANILHA.name)
        return 0
    print("4/4  Rodando a rotina de terça...\n")
    codigo = subprocess.run([sys.executable, "rotina_terca.py"], cwd=str(cfg.BASE)).returncode
    if codigo != 0:
        return codigo
    limpar_metadados()
    print("\nPronto: %s" % cfg.PLANILHA.name)
    print("Abra a planilha, habilite o conteúdo (macros) e use os botões da aba Conferência NF.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
